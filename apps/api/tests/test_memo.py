import io

from docx import Document

from casewise_api.memo import MemoIn, build_memo
from casewise_api.quotecheck import QuoteCheck
from casewise_api.rag import Source

MD = """# Notice Requirements

**Short answer:** 30, 60, or 90 days [S1].

## Details

- Under one year: *30 days* [S1]
- Over two years: 90 days [S1]

| Tenancy | Notice |
|---|---|
| < 1 year | 30 days |
| > 2 years | 90 days |
"""


def _build():
    data = MemoIn(
        question="How much notice before a rent increase?",
        answer=MD,
        sources=[
            Source(marker="S1", citation="RPP § 226-C", heading="Notice of rent increase",
                   url="https://example.com", source_type="statute"),
        ],
        quote_check=QuoteCheck(total=2, verified=1, unverified=["made up words"]),
        model="claude-opus-5-5",
    )
    return Document(io.BytesIO(build_memo(data)))


def test_memo_structure():
    doc = _build()
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert "Research Memorandum" in headings
    assert "Question Presented" in headings
    assert "Table of Authorities" in headings
    assert "Verification Appendix" in headings


def test_memo_content_rendered():
    doc = _build()
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "RPP § 226-C" in text
    assert "30, 60, or 90 days [S1]." in text
    assert "made up words" in text  # unverified quote surfaced in appendix
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "Tenancy"
